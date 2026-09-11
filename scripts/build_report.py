# -*- coding: utf-8 -*-
import json
import sys
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent
sys.path.insert(0, str(SCRIPT_DIR))
from teachers import teachers

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

data = json.load(open(REPO_ROOT / "data" / "results.json", encoding="utf-8"))

DAY_ORDER = {
    "Понедельник": 1, "Вторник": 2, "Среда": 3, "Четверг": 4,
    "Пятница": 5, "Суббота": 6, "Воскресенье": 7,
}

def clean_day(d):
    if not d:
        return ""
    # formats like "Дүйсенбі/Понедельник" or "Понедельник/Дүйсенбі" or "Среда/Сәрсенбі"
    parts = [p.strip() for p in d.split("/")]
    for p in parts:
        for ru in DAY_ORDER:
            if ru.lower() == p.lower():
                return ru
    return d

def day_sort_key(d):
    ru = clean_day(d)
    return DAY_ORDER.get(ru, 99)

def time_sort_key(t):
    if not t:
        return (99, 99)
    t = t.strip()
    try:
        h, m = t.split("-")[0].split(".")
        return (int(h), int(m))
    except Exception:
        return (99, 99)

for d in data:
    d["День_чист"] = clean_day(d["День"])

data.sort(key=lambda d: (d["Преподаватель"], day_sort_key(d["День"]), time_sort_key(d["Время"])))

NOT_FOUND = [t for t in teachers if t not in {d["Преподаватель"] for d in data}]

wb = openpyxl.Workbook()

# ---------- Sheet 1: Summary ----------
ws1 = wb.active
ws1.title = "Сводка"
font_title = Font(name="Arial", size=14, bold=True)
font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
font_normal = Font(name="Arial", size=10)
fill_header = PatternFill("solid", fgColor="305496")
fill_missing = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws1["A1"] = "Кафедра ИС — занятия преподавателей (2026-2027 уч. год, 1 акад. период)"
ws1["A1"].font = font_title
ws1.merge_cells("A1:D1")

ws1["A3"] = "ФИО преподавателя"
ws1["B3"] = "Кол-во найденных занятий"
ws1["C3"] = "Статус"
for c in ("A3", "B3", "C3"):
    ws1[c].font = font_header
    ws1[c].fill = fill_header
    ws1[c].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

from collections import Counter
counts = Counter(d["Преподаватель"] for d in data)

row = 4
for t in teachers:
    ws1.cell(row=row, column=1, value=t).font = font_normal
    n = counts.get(t, 0)
    ws1.cell(row=row, column=2, value=n).font = font_normal
    ws1.cell(row=row, column=2).alignment = Alignment(horizontal="center")
    status = "найдено" if n > 0 else "не найдено в загруженных файлах"
    cell_status = ws1.cell(row=row, column=3, value=status)
    cell_status.font = font_normal
    if n == 0:
        for col in (1, 2, 3):
            ws1.cell(row=row, column=col).fill = fill_missing
    for col in (1, 2, 3):
        ws1.cell(row=row, column=col).border = border
    row += 1

ws1.column_dimensions["A"].width = 42
ws1.column_dimensions["B"].width = 22
ws1.column_dimensions["C"].width = 34
ws1.freeze_panes = "A4"

note_row = row + 1
ws1.cell(row=note_row, column=1,
          value="Источники: файлы расписаний 1–4 курсов (бакалавриат), магистратуры (проф. и научно-пед. направления), докторантуры на 2026-2027 уч. год, 1 акад. период.").font = Font(name="Arial", size=9, italic=True)
ws1.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=3)

# ---------- Sheet 2: All classes ----------
ws2 = wb.create_sheet("Все занятия")
headers = ["ФИО преподавателя", "Уровень/Курс", "День", "Время", "Группа/Поток", "Занятие", "Файл-источник", "Лист"]
for i, h in enumerate(headers, start=1):
    c = ws2.cell(row=1, column=i, value=h)
    c.font = font_header
    c.fill = fill_header
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for r, d in enumerate(data, start=2):
    ws2.cell(row=r, column=1, value=d["Преподаватель"]).font = font_normal
    ws2.cell(row=r, column=2, value=d["Уровень/Курс"]).font = font_normal
    ws2.cell(row=r, column=3, value=d["День_чист"]).font = font_normal
    ws2.cell(row=r, column=4, value=d["Время"]).font = font_normal
    ws2.cell(row=r, column=5, value=d["Группа/Поток"]).font = font_normal
    cell_lesson = ws2.cell(row=r, column=6, value=d["Занятие"])
    cell_lesson.font = font_normal
    cell_lesson.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.cell(row=r, column=7, value=d["Файл"]).font = Font(name="Arial", size=8, color="808080")
    ws2.cell(row=r, column=8, value=d["Лист"]).font = Font(name="Arial", size=8, color="808080")
    for col in range(1, 9):
        ws2.cell(row=r, column=col).border = border

widths = [30, 26, 12, 12, 26, 60, 34, 20]
for i, w in enumerate(widths, start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"
ws2.auto_filter.ref = f"A1:H{len(data)+1}"

# ---------- Sheet 3: One sheet per teacher would be too many; instead group visually ----------
# Add a per-teacher pivot-like sheet: "По преподавателям" grouping with subtotal headers
ws3 = wb.create_sheet("По преподавателям")
r = 1
for t in teachers:
    rows_t = [d for d in data if d["Преподаватель"] == t]
    cell = ws3.cell(row=r, column=1, value=f"{t}  ({len(rows_t)} занятий)" if rows_t else f"{t}  — не найдено в загруженных расписаниях")
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    ws3.cell(row=r, column=1).fill = fill_header if rows_t else PatternFill("solid", fgColor="BF9000")
    ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    r += 1
    if rows_t:
        sub_headers = ["Уровень/Курс", "День", "Время", "Группа/Поток", "Занятие", "Файл-источник"]
        for i, h in enumerate(sub_headers, start=1):
            c = ws3.cell(row=r, column=i, value=h)
            c.font = Font(name="Arial", size=9, bold=True)
            c.fill = PatternFill("solid", fgColor="D9E1F2")
        r += 1
        for d in rows_t:
            ws3.cell(row=r, column=1, value=d["Уровень/Курс"]).font = font_normal
            ws3.cell(row=r, column=2, value=d["День_чист"]).font = font_normal
            ws3.cell(row=r, column=3, value=d["Время"]).font = font_normal
            ws3.cell(row=r, column=4, value=d["Группа/Поток"]).font = font_normal
            wc = ws3.cell(row=r, column=5, value=d["Занятие"])
            wc.font = font_normal
            wc.alignment = Alignment(wrap_text=True, vertical="top")
            ws3.cell(row=r, column=6, value=d["Файл"]).font = Font(name="Arial", size=8, color="808080")
            for col in range(1, 7):
                ws3.cell(row=r, column=col).border = border
            r += 1
    r += 1  # blank row between teachers

widths3 = [24, 12, 12, 26, 55, 34]
for i, w in enumerate(widths3, start=1):
    ws3.column_dimensions[get_column_letter(i)].width = w
ws3.freeze_panes = "A1"

out_path = REPO_ROOT / "dist" / "Кафедра_ИС_занятия_преподавателей_2026-2027.xlsx"
wb.save(out_path)
print("Saved:", out_path)
print("Total rows in 'Все занятия':", len(data))
print("Not found:", NOT_FOUND)
